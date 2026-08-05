"""Build Item Master + Opening Stock import pack from Assets office detail.xlsx."""

from __future__ import annotations

import collections
import csv
import re
from pathlib import Path

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

SRC = Path(r"c:\Users\ssapkale\Documents\Assets office detail.xlsx")
OUT_DIR = Path(__file__).resolve().parent


def clean(v) -> str:
    if v is None:
        return ""
    s = str(v).replace("\n", " ").strip()
    return re.sub(r"\s+", " ", s)


def norm_make(m: str) -> str:
    m = clean(m).upper()
    aliases = {
        "DELL": "DELL",
        "HP": "HP",
        "LENOVO": "LENOVO",
        "ACER": "ACER",
        "ASUS": "ASUS",
        "TOSHIBA": "TOSHIBA",
        "SAMSUNG": "SAMSUNG",
        "AOC": "AOC",
        "APPLE": "APPLE",
        "BAUMER": "BAUMER",
        "NOC": "NOC",
        "COMPAQ": "HP",
        "LOGITECH": "LOGITECH",
    }
    for key, val in aliases.items():
        if m == key or m.startswith(key + " ") or m.startswith(key + "-"):
            return val
    return m or "UNKNOWN"


def slug(s: str, n: int = 24) -> str:
    s = re.sub(r"[^A-Za-z0-9]+", "-", clean(s).upper()).strip("-")
    return s[:n]


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.load_workbook(SRC, data_only=True)

    catalog: dict[tuple, dict] = {}
    units: list[dict] = []

    def add_unit(
        device,
        make,
        model,
        serial,
        ram="",
        storage="",
        user="",
        remark="",
        source="",
        monitor_make="",
        monitor_model="",
        accessories="",
        product_no="",
        ip="",
        mac="",
    ):
        device = clean(device).upper() or "ASSET"
        if device in ("LAPTOP", "NOTEBOOK") or "LAPTOP" in device:
            device = "LAPTOP"
        elif device.startswith("DESK") or "DESKTOP" in device:
            device = "DESKTOP"
        make_n = norm_make(make)
        model_c = clean(model)
        serial_c = clean(serial).replace(" ", "")
        if not serial_c and not model_c:
            return
        if serial_c.upper() in ("SERIAL NO.", "SERIALNUMBER", "SR.NO", "SERIAL"):
            return
        if make_n in ("MAKE", "CATEGORY"):
            return

        if device == "LAPTOP":
            sub, prefix = "SC-LAPTOP", "LAP"
        elif device == "DESKTOP":
            sub, prefix = "SC-DESKTOP", "DESK"
        elif "CAMERA" in device or "BEAGLE" in device:
            sub, prefix = "SC-PERIPH", "CAM"
        elif "MONITOR" in device:
            sub, prefix = "SC-PERIPH", "MON"
        else:
            sub, prefix = "SC-PERIPH", "IT"

        name = f"{make_n} {model_c}".strip() or device.title()
        key = (device, make_n, re.sub(r"\s+", " ", model_c.upper()))
        if key not in catalog:
            code = f"{prefix}-{slug(make_n, 8)}-{slug(model_c, 18)}"
            base = code
            i = 2
            existing = {c["itemCode"] for c in catalog.values()}
            while code in existing:
                code = f"{base}-{i}"
                i += 1
            catalog[key] = {
                "itemCode": code,
                "itemName": name[:150],
                "itemType": "asset",
                "categoryCode": "CAT-IT",
                "subcategoryCode": sub,
                "uomCode": "NOS",
                "makeBrand": make_n.title() if make_n != "UNKNOWN" else "",
                "model": model_c[:100],
                "ram": clean(ram)[:50],
                "storage": clean(storage)[:100],
                "isSerialized": "Y",
                "isActive": "Y",
                "description": f"{device} catalog from office assets workbook",
            }
        else:
            if clean(ram) and not catalog[key]["ram"]:
                catalog[key]["ram"] = clean(ram)[:50]
            if clean(storage) and not catalog[key]["storage"]:
                catalog[key]["storage"] = clean(storage)[:100]

        status_note = clean(user)
        rem = clean(remark)
        joined = (status_note + " " + rem).lower()
        if any(x in joined for x in ["instock", "in stock", "custody", "custdy", "available"]):
            stock_hint = "IN_STORE"
        elif status_note:
            stock_hint = "ASSIGNED"
        else:
            stock_hint = "IN_STORE"

        units.append(
            {
                "sourceSheet": source,
                "itemCode": catalog[key]["itemCode"],
                "itemName": catalog[key]["itemName"],
                "itemType": "asset",
                "qty": 1,
                "uomCode": "NOS",
                "serialNo": serial_c[:100],
                "makeBrand": catalog[key]["makeBrand"],
                "model": model_c,
                "ram": clean(ram),
                "storage": clean(storage),
                "productNo": clean(product_no),
                "assignedToName": status_note if stock_hint == "ASSIGNED" else "",
                "stockStatusHint": stock_hint,
                "ipAddress": clean(ip),
                "macAddress": clean(mac),
                "monitorMake": clean(monitor_make),
                "monitorModel": clean(monitor_model),
                "accessories": clean(accessories)[:200],
                "remark": rem[:200],
            }
        )

    # Laptop Detail
    ws = wb["Laptop Detail"]
    for i, row in enumerate(ws.iter_rows(values_only=True), 1):
        if i < 4:
            continue
        vals = list(row) + [None] * 12
        (
            _sr,
            _cat,
            make,
            model,
            serial,
            warranty,
            ram,
            hdd,
            user,
            mon_make,
            mon_model,
            kb,
        ) = vals[:12]
        if not clean(make) and not clean(serial):
            continue
        add_unit(
            "LAPTOP",
            make,
            model,
            serial,
            ram,
            hdd,
            user,
            warranty or "",
            "Laptop Detail",
            mon_make,
            mon_model,
            kb,
        )

    # Desktop detail
    ws = wb["Desktop detail"]
    for i, row in enumerate(ws.iter_rows(values_only=True), 1):
        if i < 3:
            continue
        vals = list(row) + [None] * 12
        (
            _sr,
            _cat,
            make,
            model,
            serial,
            warranty,
            ram,
            hdd,
            user,
            mon_make,
            mon_model,
            kb,
        ) = vals[:12]
        if not clean(make) and not clean(serial):
            continue
        add_unit(
            "DESKTOP",
            make,
            model,
            serial,
            ram,
            hdd,
            user,
            warranty or "",
            "Desktop detail",
            mon_make,
            mon_model,
            kb,
        )

    # Available Laptop
    ws = wb["Available Laptop"]
    for i, row in enumerate(ws.iter_rows(values_only=True), 1):
        if i < 3:
            continue
        vals = list(row) + [None] * 10
        _sr, device, make, model, serial, _w, ram, storage, last_user, remark = vals[:10]
        if not clean(make) and not clean(serial):
            continue
        add_unit(
            device or "LAPTOP",
            make,
            model,
            serial,
            ram,
            storage,
            last_user,
            remark or "Instock",
            "Available Laptop",
        )

    # Available Desktop
    ws = wb["Available Desktop"]
    for i, row in enumerate(ws.iter_rows(values_only=True), 1):
        if i < 2:
            continue
        vals = [clean(c) for c in row]
        if not vals or not vals[0]:
            continue
        if vals[0].upper().startswith("DESK"):
            add_unit(
                "DESKTOP",
                vals[1] if len(vals) > 1 else "",
                vals[2] if len(vals) > 2 else "",
                vals[3] if len(vals) > 3 else "",
                vals[5] if len(vals) > 5 else "",
                vals[6] if len(vals) > 6 else "",
                vals[7] if len(vals) > 7 else "",
                "Available/Custody",
                "Available Desktop",
            )

    # Available laptop 1
    ws = wb[" Available laptop 1"]
    for i, row in enumerate(ws.iter_rows(values_only=True), 1):
        if i < 4:
            continue
        vals = [clean(c) for c in row]
        if len(vals) < 5:
            continue
        if not vals[2] and not vals[4]:
            continue
        if vals[2].upper() == "MAKE":
            continue
        add_unit(
            vals[1] or "LAPTOP",
            vals[2],
            vals[3],
            vals[4],
            vals[7] if len(vals) > 7 else "",
            vals[8] if len(vals) > 8 else "",
            vals[6] if len(vals) > 6 else "",
            (vals[10] if len(vals) > 10 else "") or (vals[9] if len(vals) > 9 else ""),
            " Available laptop 1",
            product_no=vals[5] if len(vals) > 5 else "",
        )

    # Beagle Y-AI Device
    ws = wb["Beagle Y-AI Device"]
    for i, row in enumerate(ws.iter_rows(values_only=True), 1):
        if i < 3:
            continue
        vals = list(row) + [None] * 8
        _sr, device, make, model, serial, qty, issued, remark = vals[:8]
        if not clean(make) and not clean(serial):
            continue
        if clean(device).upper() in ("DEVICE", "SR.NO"):
            continue
        try:
            q = int(float(clean(qty) or 1))
        except ValueError:
            q = 1
        serial_base = clean(serial)
        for n in range(max(q, 1)):
            serial_out = serial_base if q == 1 else (f"{serial_base}-{n + 1}" if serial_base else "")
            add_unit(
                device or "CAMERA",
                make,
                model,
                serial_out,
                "",
                "",
                issued or "",
                remark or "",
                "Beagle Y-AI Device",
            )

    # Instock Material (consumables)
    consumables: list[dict] = []
    ws = wb["Instock Material"]
    for i, row in enumerate(ws.iter_rows(values_only=True), 1):
        if i < 4:
            continue
        vals = [clean(c) for c in row]
        non_empty = [v for v in vals if v]
        if not non_empty:
            continue
        name, qty = "", ""
        if len(vals) >= 3 and vals[1] and vals[2].replace(".", "", 1).isdigit():
            name, qty = vals[1], vals[2]
        elif len(vals) >= 2 and vals[0] and vals[1].replace(".", "", 1).isdigit():
            name, qty = vals[0], vals[1]
        else:
            name = non_empty[0]
            qty = non_empty[1] if len(non_empty) > 1 else "1"
        if name.lower() in {
            "in stock material",
            "material",
            "item",
            "particular",
            "description",
            "sr.no",
            "sr no",
        }:
            continue
        try:
            qn = float(qty) if qty else 1.0
        except ValueError:
            qn = 1.0
        code = f"CONS-{slug(name, 20)}"
        consumables.append(
            {
                "itemCode": code,
                "itemName": name[:150],
                "itemType": "consumable",
                "categoryCode": "CAT-CONS",
                "subcategoryCode": "SC-CONSGEN",
                "uomCode": "PCS",
                "makeBrand": "",
                "model": "",
                "ram": "",
                "storage": "",
                "isSerialized": "N",
                "isActive": "Y",
                "description": "From Instock Material sheet",
                "openingQty": qn,
                "sourceSheet": "Instock Material",
            }
        )

    for c in consumables:
        key = ("CONS", c["itemCode"], c["itemName"].upper())
        if key not in catalog:
            catalog[key] = {
                k: c[k]
                for k in [
                    "itemCode",
                    "itemName",
                    "itemType",
                    "categoryCode",
                    "subcategoryCode",
                    "uomCode",
                    "makeBrand",
                    "model",
                    "ram",
                    "storage",
                    "isSerialized",
                    "isActive",
                    "description",
                ]
            }

    seen_serial: set[str] = set()
    unique_units: list[dict] = []
    dup = 0
    blank_serial = 0
    for u in units:
        s = u["serialNo"].upper()
        if not s:
            blank_serial += 1
            unique_units.append(u)
            continue
        if s in seen_serial:
            dup += 1
            continue
        seen_serial.add(s)
        unique_units.append(u)

    items = sorted(catalog.values(), key=lambda x: (x["itemType"], x["itemCode"]))

    out = Workbook()
    thin = Border(
        left=Side(style="thin", color="B0B8C8"),
        right=Side(style="thin", color="B0B8C8"),
        top=Side(style="thin", color="B0B8C8"),
        bottom=Side(style="thin", color="B0B8C8"),
    )
    hdr_fill = PatternFill("solid", fgColor="0EA5E9")
    hdr_font = Font(bold=True, color="FFFFFF")

    def write_sheet(ws, headers, rows, widths=None):
        for c, h in enumerate(headers, 1):
            cell = ws.cell(1, c, h)
            cell.fill = hdr_fill
            cell.font = hdr_font
            cell.border = thin
        for r, row in enumerate(rows, 2):
            for c, h in enumerate(headers, 1):
                cell = ws.cell(r, c, row.get(h, ""))
                cell.border = thin
        if widths:
            for i, w in enumerate(widths, 1):
                ws.column_dimensions[get_column_letter(i)].width = w

    ws = out.active
    ws.title = "README"
    notes = [
        "CAITS Import Pack — Office Assets (from Assets office detail.xlsx)",
        "",
        "How to use",
        "1. Verify Category CAT-IT / CAT-CONS and Subcategories SC-LAPTOP, SC-DESKTOP, SC-PERIPH, SC-CONSGEN.",
        "2. Ensure Unit Master has NOS (assets) and PCS (consumables).",
        "3. Create Item Master rows from sheet ITEM_MASTER first (one catalog row per Make+Model).",
        "4. Post Opening Stock from OPENING_STOCK_IMPORT: Item Type = Asset, Qty = 1 per serial line.",
        "5. Post consumables from CONSUMABLES_OPENING with Qty as given (Item Type = Consumable).",
        "6. assignedToName / stockStatusHint are reference only — map employee & location in CAITS after import.",
        "7. Duplicate serials across sheets were de-duplicated; blank serial rows need review before posting.",
        "",
        "Source sheets used",
        "Laptop Detail, Desktop detail, Available Laptop, Available Desktop, Available laptop 1, Beagle Y-AI Device, Instock Material",
        "",
        "Not used for stock import (network / issue logs)",
        "IT Park office IP update, Pharma office IP update, Old laptop Detail, Tempeally Issued Material, Return Issued Material",
        "",
        "Summary",
        f"Item Master catalog rows: {len(items)}",
        f"Opening stock asset unit lines: {len(unique_units)}",
        f"Duplicate serials skipped: {dup}",
        f"Rows with blank serial: {blank_serial}",
        f"Consumable / instock material lines: {len(consumables)}",
    ]
    for r, text in enumerate(notes, 1):
        ws.cell(r, 1, text)
    ws.column_dimensions["A"].width = 110

    item_headers = [
        "itemCode",
        "itemName",
        "itemType",
        "categoryCode",
        "subcategoryCode",
        "uomCode",
        "makeBrand",
        "model",
        "ram",
        "storage",
        "isSerialized",
        "isActive",
        "description",
    ]
    write_sheet(
        out.create_sheet("ITEM_MASTER"),
        item_headers,
        items,
        [18, 36, 12, 12, 16, 10, 14, 28, 10, 16, 12, 10, 40],
    )

    ost_headers = [
        "itemCode",
        "itemName",
        "itemType",
        "qty",
        "uomCode",
        "serialNo",
        "ram",
        "storage",
        "productNo",
        "ipAddress",
        "macAddress",
        "assignedToName",
        "stockStatusHint",
        "monitorMake",
        "monitorModel",
        "accessories",
        "remark",
        "sourceSheet",
    ]
    write_sheet(
        out.create_sheet("OPENING_STOCK_IMPORT"),
        ost_headers,
        unique_units,
        [18, 28, 10, 6, 8, 18, 10, 14, 14, 14, 16, 24, 14, 12, 16, 24, 24, 18],
    )

    cons_headers = [
        "itemCode",
        "itemName",
        "itemType",
        "qty",
        "uomCode",
        "remark",
        "sourceSheet",
    ]
    cons_rows = [
        {
            "itemCode": c["itemCode"],
            "itemName": c["itemName"],
            "itemType": "consumable",
            "qty": c["openingQty"],
            "uomCode": "PCS",
            "remark": c["description"],
            "sourceSheet": c["sourceSheet"],
        }
        for c in consumables
    ]
    write_sheet(
        out.create_sheet("CONSUMABLES_OPENING"),
        cons_headers,
        cons_rows,
        [18, 40, 12, 8, 8, 30, 18],
    )

    counts = collections.Counter(u["itemCode"] for u in unique_units)
    code_to_item = {i["itemCode"]: i for i in items}
    sum_rows = []
    for code, cnt in sorted(counts.items(), key=lambda x: -x[1]):
        it = code_to_item.get(code, {})
        sum_rows.append(
            {
                "itemCode": code,
                "itemName": it.get("itemName", ""),
                "makeBrand": it.get("makeBrand", ""),
                "model": it.get("model", ""),
                "unitCount": cnt,
            }
        )
    write_sheet(
        out.create_sheet("SUMMARY_BY_MODEL"),
        ["itemCode", "itemName", "makeBrand", "model", "unitCount"],
        sum_rows,
        [18, 36, 14, 28, 10],
    )

    out_path = OUT_DIR / "CAITS_Office_Assets_Import.xlsx"
    out.save(out_path)

    with open(OUT_DIR / "ITEM_MASTER.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=item_headers)
        w.writeheader()
        w.writerows(items)
    with open(OUT_DIR / "OPENING_STOCK_IMPORT.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=ost_headers, extrasaction="ignore")
        w.writeheader()
        w.writerows(unique_units)

    print("OUT", out_path)
    print(
        "items",
        len(items),
        "units",
        len(unique_units),
        "dup",
        dup,
        "blank_serial",
        blank_serial,
        "consumables",
        len(consumables),
    )
    print("top models:")
    for r in sum_rows[:15]:
        print(r["unitCount"], r["itemCode"], r["itemName"])


if __name__ == "__main__":
    main()
